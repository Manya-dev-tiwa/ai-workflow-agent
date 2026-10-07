import openpyxl
from copy import copy

path = "data/workflows.xlsx"
wb = openpyxl.load_workbook(path)
ws = wb["Test_Questions"]

NEW_WF009 = (
    "Assign this urgent task to the best available developer: "
    "fix the payment API timeout bug in our Python FastAPI service."
)
WF011_ROW = (
    "WF011",
    "Which orders are overdue?",
    "Tests date comparison and filtering",
)

# 1. Update the WF009 question (match on column A)
updated = False
for row in ws.iter_rows(min_row=2):
    if row[0].value == "WF009":
        row[1].value = NEW_WF009
        updated = True
        break
print("WF009 updated" if updated else "WF009 row NOT found")

# 2. Add WF011 if it is not already there
ids = [r[0].value for r in ws.iter_rows(min_row=2)]
if "WF011" in ids:
    print("WF011 already present, skipped")
else:
    last = max(r[0].row for r in ws.iter_rows(min_row=2) if r[0].value)
    new_row = last + 1
    for col, value in enumerate(WF011_ROW, start=1):
        cell = ws.cell(row=new_row, column=col, value=value)
        src = ws.cell(row=last, column=col)
        cell.font = copy(src.font)
        cell.alignment = copy(src.alignment)
        cell.border = copy(src.border)
    print(f"WF011 added at row {new_row}")

wb.save(path)
print("Saved")